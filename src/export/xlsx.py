from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import OLLAMA_MODEL
from src.export.markdown import parse_markdown_table

HDR_FILL = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
HDR_FONT = Font(color="1E40AF", bold=True)
ALT_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
ALT_FONT = Font(color="1E293B")
WRAP = Alignment(wrap_text=True, vertical="top")


def _set_header(ws, col: int, row: int, value: str):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = HDR_FILL
    c.font = HDR_FONT
    c.alignment = WRAP


def build_xlsx(vals: dict) -> BytesIO:
    """Genera el Excel (artículos, matriz y metadatos de búsqueda) a partir del estado de la sesión."""
    papers = vals.get("papers", [])
    matrix_md = vals.get("matrix", "")
    equation = vals.get("search_equation", "")
    needs = vals.get("research_needs", "")
    paper_count = vals.get("paper_count", 5)

    wb = openpyxl.Workbook()

    # Sheet 1: Artículos
    ws1 = wb.active
    ws1.title = "Artículos"

    cols1 = [
        ("#", 4), ("Título", 48), ("Autores", 28), ("Año", 6),
        ("Fuente", 14), ("Revista / Conferencia", 30), ("DOI", 24),
        ("Citas", 8), ("Open Access", 12), ("Palabras Clave", 28),
        ("Volumen", 8), ("Número", 8), ("Páginas", 10), ("URL", 50),
    ]
    for i, (header, width) in enumerate(cols1, 1):
        _set_header(ws1, i, 1, header)
        ws1.column_dimensions[get_column_letter(i)].width = width

    for i, p in enumerate(papers, 1):
        row = i + 1
        values = [
            i, p.get("title", ""), p.get("authors", ""), p.get("year", ""),
            p.get("source", ""), p.get("journal", ""), p.get("doi", ""),
            p.get("citations", ""), p.get("open_access", ""), p.get("keywords", ""),
            p.get("volume", ""), p.get("issue", ""), p.get("pages", ""), p.get("url", ""),
        ]
        for col, val in enumerate(values, 1):
            cell = ws1.cell(row=row, column=col, value=val)
            cell.alignment = WRAP
            if i % 2 == 0:
                cell.fill = ALT_FILL
                cell.font = ALT_FONT

    # Sheet 2: Matriz
    ws2 = wb.create_sheet("Matriz Bibliográfica")
    if matrix_md:
        rows = parse_markdown_table(matrix_md)
        for r, row_data in enumerate(rows, 1):
            for c, val in enumerate(row_data, 1):
                cell = ws2.cell(row=r, column=c, value=val)
                cell.alignment = WRAP
                if r == 1:
                    cell.fill = HDR_FILL
                    cell.font = HDR_FONT
                elif r % 2 == 0:
                    cell.fill = ALT_FILL
                    cell.font = ALT_FONT
        max_col = max((len(r) for r in rows), default=1)
        for c in range(1, max_col + 1):
            ws2.column_dimensions[get_column_letter(c)].width = 28

    # Sheet 3: Búsqueda
    ws3 = wb.create_sheet("Búsqueda")
    meta = [
        ("Tema de investigación", needs),
        ("Ecuación de búsqueda", equation),
        ("Papers por fuente", paper_count),
        ("Total artículos", len(papers)),
        ("ArXiv", sum(1 for p in papers if p.get("source") == "ArXiv")),
        ("Google Scholar", sum(1 for p in papers if p.get("source") == "Google Scholar")),
        ("Subidos", sum(1 for p in papers if p.get("source") == "Subido")),
        ("Fecha de búsqueda", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Modelo LLM", f"{OLLAMA_MODEL} (Ollama)"),
    ]
    for r, (k, v) in enumerate(meta, 1):
        ws3.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=r, column=2, value=str(v)).alignment = WRAP
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 65

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
