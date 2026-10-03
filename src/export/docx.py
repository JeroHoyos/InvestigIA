import re
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.export.markdown import parse_markdown_table


def build_docx(vals: dict) -> BytesIO:
    """Genera el reporte Word con la metodología, artículos y matriz a partir del estado de la sesión."""
    papers = vals.get("papers", [])
    matrix_md = vals.get("matrix", "")
    equation = vals.get("search_equation", "")
    explanation = vals.get("equation_explanation", "")
    needs = vals.get("research_needs", "")
    matrix_template = vals.get("matrix_template", "")

    doc = Document()

    # Styles
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def set_heading(para, text, level=1):
        para.style = f"Heading {level}"
        para.runs[0].font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
        return para

    def add_heading(doc, text, level=1):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
        return h

    # Cover
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("Reporte de Investigación Bibliográfica")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run("Science Agent — DataHack 2026")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.add_run(f"Generado: {datetime.now().strftime('%d de %B de %Y, %H:%M')}")

    doc.add_page_break()

    # Section 1: Research topic
    add_heading(doc, "1. Tema de Investigación", 1)
    doc.add_paragraph(needs or "No especificado.")

    # Section 2: Methodology
    add_heading(doc, "2. Metodología", 1)
    doc.add_paragraph(
        "Este reporte fue generado utilizando Science Agent, un sistema de investigación "
        "bibliográfica basado en inteligencia artificial. El proceso siguió los siguientes pasos:"
    )

    steps = [
        ("Definición del tema", "El investigador describió su área de investigación o pregunta científica al agente conversacional."),
        ("Generación de ecuación de búsqueda", "Un modelo de lenguaje large (LLM) local mediante Ollama ({OLLAMA_MODEL}) construyó una ecuación booleana elaborada con sinónimos, variantes morfológicas y operadores de campo específicos."),
        ("Búsqueda paralela", "El sistema consultó simultáneamente ArXiv (repositorio de preprints científicos) y Google Scholar, recuperando artículos con sus metadatos completos: título, autores, año, abstract, DOI, citas, revista y acceso abierto."),
        ("Selección de plantilla", "El investigador eligió una plantilla de matriz bibliográfica adaptada a su tipo de revisión (Estado del Arte, Revisión Sistemática, Benchmarking, Marco Teórico, Tendencias o Meta-análisis)."),
        ("Generación de matriz", "El LLM procesó todos los artículos encontrados y generó una matriz bibliográfica estructurada en Markdown, incluyendo análisis de tendencias dominantes, brechas identificadas y recomendaciones de lectura."),
        ("Análisis y exportación", "Los resultados fueron exportados a este documento Word con la descripción metodológica completa y los metadatos de todos los artículos."),
    ]

    for i, (title, desc) in enumerate(steps, 1):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(f"{title}: ")
        run.bold = True
        p.add_run(desc)

    # Section 3: Search equation
    add_heading(doc, "3. Ecuación de Búsqueda", 1)
    doc.add_paragraph("La siguiente ecuación booleana fue utilizada para la recuperación de artículos:")

    eq_para = doc.add_paragraph()
    eq_run = eq_para.add_run(equation or "No disponible")
    eq_run.font.name = "Courier New"
    eq_run.font.size = Pt(9)
    eq_para.paragraph_format.left_indent = Cm(1)

    if explanation:
        add_heading(doc, "Explicación de la ecuación", 2)
        doc.add_paragraph(explanation)

    # Section 4: Papers
    add_heading(doc, "4. Artículos Recuperados", 1)

    arxiv_n = sum(1 for p in papers if p.get("source") == "ArXiv")
    scholar_n = sum(1 for p in papers if p.get("source") == "Google Scholar")
    uploaded_n = sum(1 for p in papers if p.get("source") == "Subido")

    doc.add_paragraph(
        f"Se recuperaron {len(papers)} artículos en total: {arxiv_n} de ArXiv, "
        f"{scholar_n} de Google Scholar y {uploaded_n} subidos manualmente."
    )

    for i, p in enumerate(papers, 1):
        add_heading(doc, f"{i}. {p.get('title', 'Sin título')}", 2)

        meta_lines = []
        if p.get("authors"):
            meta_lines.append(("Autores", p["authors"]))
        if p.get("year"):
            meta_lines.append(("Año", str(p["year"])))
        if p.get("source"):
            meta_lines.append(("Fuente", p["source"]))
        if p.get("journal"):
            meta_lines.append(("Revista", p["journal"]))
        if p.get("doi"):
            meta_lines.append(("DOI", p["doi"]))
        if p.get("citations"):
            meta_lines.append(("Citas", str(p["citations"])))
        if p.get("open_access"):
            meta_lines.append(("Acceso Abierto", p["open_access"]))
        if p.get("url"):
            meta_lines.append(("URL", p["url"]))

        for label, value in meta_lines:
            p_meta = doc.add_paragraph()
            run_l = p_meta.add_run(f"{label}: ")
            run_l.bold = True
            p_meta.add_run(value)

        if p.get("abstract"):
            doc.add_paragraph(f"Resumen: {p['abstract'][:600]}{'...' if len(p.get('abstract','')) > 600 else ''}")

        if p.get("keywords"):
            kw_p = doc.add_paragraph()
            kw_r = kw_p.add_run("Palabras clave: ")
            kw_r.bold = True
            kw_p.add_run(p["keywords"])

    # Section 5: Matrix
    if matrix_md:
        doc.add_page_break()
        add_heading(doc, "5. Matriz Bibliográfica", 1)

        template_names = {
            "estado_arte": "Estado del Arte",
            "sistematica": "Revisión Sistemática",
            "benchmarking": "Benchmarking Técnico",
            "marco_teorico": "Marco Teórico",
            "tendencias": "Tendencias e Innovación",
            "metaanalisis": "Meta-análisis",
        }
        tpl_name = template_names.get(matrix_template, matrix_template or "Personalizada")
        doc.add_paragraph(f"Plantilla utilizada: {tpl_name}")

        rows = parse_markdown_table(matrix_md)
        if rows:
            try:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r_idx, row_data in enumerate(rows):
                    for c_idx, cell_val in enumerate(row_data):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_val
                        if r_idx == 0:
                            for par in cell.paragraphs:
                                for run in par.runs:
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            tc = cell._tc
                            tcPr = tc.get_or_add_tcPr()
                            shd = OxmlElement("w:shd")
                            shd.set(qn("w:fill"), "1D4ED8")
                            shd.set(qn("w:color"), "auto")
                            shd.set(qn("w:val"), "clear")
                            tcPr.append(shd)
            except Exception:
                doc.add_paragraph("(La tabla no pudo ser generada en este formato)")
                doc.add_paragraph(matrix_md[:3000])

        # Extract trend analysis section
        trend_match = re.search(r'##\s*Análisis General(.*?)(?=^##|\Z)', matrix_md, re.DOTALL | re.MULTILINE)
        if trend_match:
            add_heading(doc, "6. Análisis de Tendencias", 1)
            trend_text = trend_match.group(1).strip()
            for line in trend_text.split("\n"):
                line = line.strip()
                if line.startswith("###"):
                    add_heading(doc, line.lstrip("#").strip(), 2)
                elif line.startswith("**") and line.endswith("**"):
                    p = doc.add_paragraph()
                    p.add_run(line.strip("*")).bold = True
                elif line.startswith("-") or line.startswith("*"):
                    doc.add_paragraph(line.lstrip("-* "), style="List Bullet")
                elif line:
                    doc.add_paragraph(line)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
